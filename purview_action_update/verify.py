"""Post-write verification: re-read the produced workbook and check it.

The planner refuses bad writes one row at a time. This checks the FILE that
came out, independently, against the same rules — the difference between
"each decision looked right" and "the artifact a human is about to upload is
valid". A bug in the writing layer, a column index slip, or a sheet clobbered
on save would all pass the planner and fail here.

Every check is either a rule transcribed from the workbook's own
"How To Update Actions" tab, or an invariant this tool promises (it touches
three columns, or five under --on-date-conflict, and nothing else).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl

from config import (
    COL_ACTION_ID,
    COL_ACTION_NAME,
    COL_IMPL_DATE,
    COL_IMPL_STATUS,
    COL_TEST_DATE,
    COL_TEST_STATUS,
    COL_TESTING_TYPE,
    PURVIEW_STATUSES,
    SHEET,
    TEST_STATUS_ALLOWED,
)
from mapping import parse_purview_datetime


def verify(original: Path, produced: Path, *, writable: tuple[str, ...]) -> list[dict[str, Any]]:
    """Return a list of violations. Empty means the workbook is safe to upload."""
    problems: list[dict[str, Any]] = []
    before = openpyxl.load_workbook(original, data_only=True)
    after = openpyxl.load_workbook(produced, data_only=True)

    def fail(check: str, detail: str, row: int | None = None) -> None:
        problems.append({"check": check, "row": row, "detail": detail})

    # --- structure ------------------------------------------------------- #
    if before.sheetnames != after.sheetnames:
        fail("sheets", f"{before.sheetnames} became {after.sheetnames}")
        return problems

    for name in before.sheetnames:
        a, b = before[name], after[name]
        if (a.max_row, a.max_column) != (b.max_row, b.max_column):
            fail("dimensions", f"{name}: {a.max_row}x{a.max_column} became "
                               f"{b.max_row}x{b.max_column}")
        if name == SHEET:
            continue
        # Reference tabs must survive the round-trip untouched.
        for row in range(1, min(a.max_row, b.max_row) + 1):
            for col in range(1, min(a.max_column, b.max_column) + 1):
                if a.cell(row, col).value != b.cell(row, col).value:
                    fail("untouched-sheet", f"{name} r{row}c{col} changed", row)
                    break

    sheet_a, sheet_b = before[SHEET], after[SHEET]
    headers = {str(c.value).strip(): c.column for c in sheet_b[1] if c.value}
    if [c.value for c in sheet_a[1]] != [c.value for c in sheet_b[1]]:
        fail("headers", "the Action Update header row changed")

    # --- per row ---------------------------------------------------------- #
    for row in range(2, sheet_b.max_row + 1):
        def old(col: str) -> Any:
            return sheet_a.cell(row, headers[col]).value

        def new(col: str) -> Any:
            return sheet_b.cell(row, headers[col]).value

        # identity and columns this tool must never touch
        for col in (COL_ACTION_ID, COL_ACTION_NAME, COL_TESTING_TYPE):
            if old(col) != new(col):
                fail("identity-changed", f"{col} was modified", row)
        for col, index in headers.items():
            if col not in writable and sheet_a.cell(row, index).value != sheet_b.cell(row, index).value:
                fail("read-only-column", f"{col} was modified", row)

        status = (str(new(COL_IMPL_STATUS)).strip() if new(COL_IMPL_STATUS) else "")
        test_status = (str(new(COL_TEST_STATUS)).strip() if new(COL_TEST_STATUS) else "")

        if status and status not in PURVIEW_STATUSES:
            fail("status-vocabulary", f"Implementation Status {status!r} is not a "
                                      f"Purview value", row)

        if status and test_status:
            allowed = TEST_STATUS_ALLOWED.get(status, set())
            if test_status not in allowed:
                fail("status-combination",
                     f"{status!r} does not permit Test Status {test_status!r} "
                     f"(allowed: {sorted(allowed)})", row)

        impl_date, test_date = new(COL_IMPL_DATE), new(COL_TEST_DATE)
        for col, value in ((COL_IMPL_DATE, impl_date), (COL_TEST_DATE, test_date)):
            if value not in (None, "") and parse_purview_datetime(value) is None:
                fail("date-format", f"{col} {value!r} is not MM/DD/YYYY HH:MM:SS", row)

        parsed_impl, parsed_test = (parse_purview_datetime(impl_date),
                                    parse_purview_datetime(test_date))
        if parsed_impl and parsed_test and parsed_test < parsed_impl:
            fail("test-date-order",
                 f"Test Date {test_date!r} is before Implementation Date "
                 f"{impl_date!r}", row)

    return problems


def summarize(problems: list[dict[str, Any]]) -> str:
    if not problems:
        return "  verification         PASS — the workbook satisfies Purview's rules"
    from collections import Counter

    counts = Counter(p["check"] for p in problems)
    lines = [f"  verification         FAIL — {len(problems)} violation(s) {dict(counts)}"]
    for problem in problems[:10]:
        where = f"row {problem['row']}" if problem["row"] else "workbook"
        lines.append(f"      [{problem['check']}] {where}: {problem['detail']}")
    if len(problems) > 10:
        lines.append(f"      … and {len(problems) - 10} more (see run_report.json)")
    return "\n".join(lines)
