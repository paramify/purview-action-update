"""Rules for the Purview Action Update writeback tool.

These tests are the specification. The tool prepares a file a human uploads
into a live compliance assessment, so the cases that matter most are the ones
where it must REFUSE to write.
"""

import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import pytest

TOOL = Path(__file__).resolve().parents[1] / "purview_action_update"
sys.path.insert(0, str(TOOL))

#: The client export moves around; the end-to-end tests skip if none is present.
#: The end-to-end tests need a real Purview export; they skip without one.
#: PURVIEW_EXPORT points at it, so no engagement-specific path is baked in.
WORKBOOK_CANDIDATES = tuple(
    Path(p) for p in (
        os.environ.get("PURVIEW_EXPORT", ""),
        "ExportActions.xlsx",
        Path.home() / "Downloads" / "ExportActions.xlsx",
        Path.home() / "Desktop" / "ExportActions.xlsx",
    ) if p
)


def find_workbook():
    return next((p for p in WORKBOOK_CANDIDATES if p.exists()), None)

import audit  # noqa: E402
import matching  # noqa: E402
from mapping import (  # noqa: E402
    format_purview_datetime,
    map_status,
    parse_purview_datetime,
    render_notes,
    check_test_date,
    check_test_status,
)
from planner import MODE_FILL_EMPTY, MODE_SYNC, plan_row  # noqa: E402

UTC = ZoneInfo("UTC")


# --- status mapping -------------------------------------------------------- #

@pytest.mark.parametrize(
    "paramify,purview",
    [
        ("IMPLEMENTED", "Implemented"),
        ("ALTERNATIVE_IMPLEMENTATION", "AlternativeImplementation"),
        ("NOT_IMPLEMENTED", "NotImplemented"),
        ("PLANNED", "Planned"),
        ("NOT_APPLICABLE", "NotInScope"),
    ],
)
def test_mappable_statuses(paramify, purview):
    assert map_status(paramify) == (purview, None)


@pytest.mark.parametrize("paramify", ["PARTIALLY_IMPLEMENTED", "NOT_SET", "", None, "WAT"])
def test_unmappable_statuses_are_refused_with_a_reason(paramify):
    value, reason = map_status(paramify)
    assert value is None and reason


def test_partially_implemented_is_not_silently_downgraded():
    # Purview has no partial state; both candidates force Test Status to None,
    # which would wipe a recorded pass.
    _, reason = map_status("PARTIALLY_IMPLEMENTED")
    assert "human decision" in reason


# --- the Test Status guard ------------------------------------------------- #

def test_writing_notimplemented_over_a_passed_test_is_blocked():
    blocking, advisory = check_test_status("NotImplemented", "Passed")
    assert blocking and not advisory


def test_allowed_combination_is_clean():
    assert check_test_status("Implemented", "Passed") == (None, None)


def test_blank_test_status_is_written_but_flagged():
    blocking, advisory = check_test_status("Implemented", None)
    assert blocking is None
    assert "blank" in advisory


def test_blank_test_status_is_refused_under_strict():
    blocking, _ = check_test_status("Implemented", None, strict=True)
    assert "strict" in blocking


def test_none_is_a_real_value_for_notimplemented():
    assert check_test_status("NotImplemented", "None") == (None, None)


# --- dates ----------------------------------------------------------------- #

def test_date_format_matches_the_exports_own():
    # Purview exported '6/23/2026 15:02:24' -- no zero padding on month/day/hour.
    assert format_purview_datetime(datetime(2026, 6, 23, 15, 2, 24)) == "6/23/2026 15:02:24"
    assert format_purview_datetime(datetime(2026, 1, 5, 9, 5, 3)) == "1/5/2026 9:05:03"


def test_dates_round_trip():
    assert parse_purview_datetime("6/23/2026 15:02:24") == datetime(2026, 6, 23, 15, 2, 24)
    assert parse_purview_datetime(None) is None
    assert parse_purview_datetime("not a date") is None


def test_implementation_date_after_test_date_is_refused():
    # Purview requires Test Date >= Implementation Date.
    later = datetime(2026, 8, 1, 12, 0, 0, tzinfo=timezone.utc)
    assert check_test_date(later, "6/23/2026 15:02:24")
    assert check_test_date(later, "9/1/2026 00:00:00") is None
    assert check_test_date(later, None) is None


# --- notes ----------------------------------------------------------------- #

def test_single_narrative_is_used_verbatim():
    text, names = render_notes([{"type": "PROVIDER", "name": "Capability Provider",
                                 "narrative": "The platform enforces it."}])
    assert text == "The platform enforces it."
    assert names == ["Capability Provider"]


def test_provider_narrative_sorts_first_and_is_labelled():
    text, names = render_notes([
        {"type": "CUSTOMER", "name": "Customer Responsibility", "narrative": "Customer part."},
        {"type": "PROVIDER", "name": "Capability Provider", "narrative": "Provider part."},
    ])
    assert text.startswith("Capability Provider: Provider part.")
    assert names == ["Capability Provider", "Customer Responsibility"]


def test_no_narrative_yields_none_so_the_cell_is_left_alone():
    # Returning "" would overwrite the client's own note -- the rules tab is
    # explicit that a written value replaces what is there.
    assert render_notes([]) == (None, [])
    assert render_notes([{"type": "PROVIDER", "narrative": "   "}]) == (None, [])


def test_overlong_narrative_is_truncated_to_the_cell_limit():
    text, _ = render_notes([{"type": "PROVIDER", "name": "P", "narrative": "x" * 40000}])
    assert len(text) <= 32767 and text.endswith("truncated]")


# --- the join -------------------------------------------------------------- #

def _cap(name, **kw):
    return {"id": kw.pop("id", "id-" + name[:6]), "name": name,
            "implementationStatus": kw.pop("status", "IMPLEMENTED"),
            "functions": kw.pop("functions", []), **kw}


def test_exact_names_match_and_orphans_are_reported_both_ways():
    result = matching.build(
        ["Alpha", "Beta", "Gamma"], [_cap("Alpha"), _cap("Beta"), _cap("Delta")]
    )
    assert set(result.exact) == {"Alpha", "Beta"}
    assert result.unmatched_actions == ["Gamma"]
    assert result.unmatched_capabilities == ["Delta"]


def test_case_drift_is_a_near_match_and_not_applied_by_default():
    result = matching.build(["Alpha Action"], [_cap("ALPHA ACTION")])
    assert result.exact == {}
    assert "Alpha Action" in result.near
    assert result.lookup("Alpha Action", accept_near=False) is None
    assert result.lookup("Alpha Action", accept_near=True) is not None


def test_duplicate_names_are_reported():
    result = matching.build(["Dup", "Dup"], [_cap("Dup")])
    assert result.duplicate_action_names == {"Dup": 2}


def test_two_capabilities_normalizing_alike_are_ambiguous_not_guessed():
    result = matching.build(["alpha"], [_cap("Alpha", id="a"), _cap("ALPHA", id="b")])
    assert result.ambiguous == ["alpha"]
    assert result.lookup("alpha", accept_near=True) is None


# --- audit log resolution -------------------------------------------------- #

def _event(rid, stamp, changes, rtype="SOLUTION_CAPABILITY", eid="e"):
    return {"id": eid, "timestamp": stamp, "type": "HISTORY",
            "resource": {"type": rtype, "id": rid},
            "action": {"verb": "UPDATE", "changes": changes}}


def test_explicit_field_name_resolves_by_field_strategy():
    got = audit.index_by_capability([
        _event("c1", "2026-07-14T08:30:00Z",
               [{"field": "implementationStatus", "old": "NOT_SET", "new": "IMPLEMENTED"}])
    ])
    assert got["c1"].method == "field"
    assert got["c1"].status_from_audit == "IMPLEMENTED"


def test_unnamed_change_resolves_by_value_vocabulary():
    # The documented example carries no field name, so the status vocabulary
    # is what identifies the change.
    got = audit.index_by_capability([
        _event("c1", "2026-07-14T08:30:00Z", [{"old": "NOT_SET", "new": "IMPLEMENTED"}])
    ])
    assert got["c1"].method == "value"


def test_newest_status_change_wins_regardless_of_page_order():
    got = audit.index_by_capability([
        _event("c1", "2026-01-02T09:00:00Z", [{"old": "NOT_SET", "new": "PLANNED"}], eid="old"),
        _event("c1", "2026-06-30T14:05:09Z", [{"old": "PLANNED", "new": "IMPLEMENTED"}], eid="new"),
    ])
    assert got["c1"].event_id == "new"


def test_non_capability_resources_are_ignored():
    assert audit.index_by_capability([
        _event("i1", "2026-08-01T10:00:00Z", [{"old": "Open", "new": "Closed"}], rtype="ISSUE")
    ]) == {}


def test_a_non_status_change_is_not_mistaken_for_one():
    assert audit.index_by_capability([
        _event("c1", "2026-08-01T10:00:00Z", [{"old": "Old name", "new": "New name"}])
    ]) == {}


def test_probe_reports_whether_change_entries_name_their_field():
    shape = audit.probe_shape([
        _event("c1", "2026-08-01T10:00:00Z", [{"old": "a", "new": "b"}])
    ])
    assert shape["change_entry_has_field_name"] is False
    assert shape["resource_type_counts"] == {"SOLUTION_CAPABILITY": 1}
    assert shape["events_carrying_changes"] == 1


def test_probe_diagnoses_a_freshly_copied_workspace():
    # Every event a CREATE with no change entries -- what copying a workspace
    # produces, and indistinguishable from "quiet log" without this diagnosis.
    created = [
        {"id": f"e{i}", "timestamp": "2026-09-28T10:00:00Z", "type": "HISTORY",
         "resource": {"type": "SOLUTION_CAPABILITY", "id": f"c{i}"},
         "action": {"verb": "CREATE", "changes": []}}
        for i in range(50)
    ]
    shape = audit.probe_shape(created)
    assert shape["events_scanned"] == 50
    assert shape["action_verb_counts"] == {"CREATE": 50}
    assert shape["events_carrying_changes"] == 0
    assert "recently COPIED" in shape["diagnosis"]


def test_probe_counts_every_event_not_just_the_newest_page():
    # Sampling the newest 40 is exactly what hid the copied-workspace signal.
    events = [
        {"id": f"e{i}", "timestamp": "2026-09-28T10:00:00Z", "type": "HISTORY",
         "resource": {"type": "SOLUTION_CAPABILITY", "id": "c1"},
         "action": {"verb": "CREATE", "changes": []}}
        for i in range(60)
    ] + [_event("c1", "2026-09-27T10:00:00Z", [{"old": "NOT_SET", "new": "IMPLEMENTED"}])]
    assert audit.probe_shape(events)["events_carrying_changes"] == 1
    assert audit.probe_shape(events, sample=40)["events_carrying_changes"] == 0


def test_probe_values_histograms_only_short_values():
    long_prose = "x" * 200
    events = [_event("c1", "2026-09-27T10:00:00Z",
                     [{"old": "NOT_SET", "new": "IMPLEMENTED"},
                      {"old": "", "new": long_prose}])]
    shape = audit.probe_shape(events, include_values=True)
    values = shape["solution_capability_change_values"]
    assert values == {"old=NOT_SET": 1, "new=IMPLEMENTED": 1}
    assert not any(long_prose in k for k in values)


# --- the row planner ------------------------------------------------------- #

BLANK_ROW = {"Action Id": "a1", "Improvement Action Name": "Alpha",
             "Implementation Status": None, "Implementation Date": None,
             "Implementation Notes": None, "Test Status": None, "Test Date": None}


def _plan(values, cap, resolution=None, mode=MODE_FILL_EMPTY, **kw):
    return plan_row(row_number=2, values=values, solcap=cap, resolution=resolution,
                    tzinfo=UTC, mode=mode, **kw)


def _written(plan, column):
    return next((u for u in plan.updates if u.column == column and u.written), None)


def _res(status="IMPLEMENTED", when=datetime(2026, 6, 30, 14, 5, 9, tzinfo=timezone.utc),
         method="value"):
    return audit.DateResolution(when, status, method)


def test_status_and_date_both_come_from_the_audit_activity():
    plan = _plan(BLANK_ROW, _cap("Alpha", status="NOT_SET"), _res("IMPLEMENTED"))
    assert _written(plan, "Implementation Status").new == "Implemented"
    assert _written(plan, "Implementation Date").new == "6/30/2026 14:05:09"


def test_status_is_the_audit_value_not_the_capabilitys_current_one():
    # The capability currently reports IMPLEMENTED; the latest activity was a
    # change to PLANNED. The Activity feed is the source, so Planned wins.
    plan = _plan(BLANK_ROW, _cap("Alpha", status="IMPLEMENTED"), _res("PLANNED"))
    assert _written(plan, "Implementation Status").new == "Planned"
    assert any(a["kind"] == "audit_disagrees_with_capability"
               and "earlier transition" in a["message"] for a in plan.advisories)


def test_audit_display_spelling_maps_the_same_as_the_enum():
    plan = _plan(BLANK_ROW, _cap("Alpha"), _res("Not Implemented"))
    assert _written(plan, "Implementation Status").new == "NotImplemented"


def test_no_status_change_activity_writes_neither_status_nor_date():
    plan = _plan(BLANK_ROW, _cap("Alpha", status="IMPLEMENTED"), None)
    assert _written(plan, "Implementation Status") is None
    assert _written(plan, "Implementation Date") is None
    assert sum("Activity feed" in s for s in plan.skipped) == 2


def test_fallback_writes_the_capability_status_but_never_a_date():
    plan = _plan(BLANK_ROW, _cap("Alpha", status="IMPLEMENTED"), None, fallback_solcap=True)
    assert _written(plan, "Implementation Status").new == "Implemented"
    assert _written(plan, "Implementation Date") is None
    assert any(a["kind"] == "status_from_capability_fallback"
               and "no Implementation Date" in a["message"] for a in plan.advisories)


def test_unmappable_audit_status_is_refused():
    plan = _plan(BLANK_ROW, _cap("Alpha"), _res("PARTIALLY_IMPLEMENTED"))
    assert _written(plan, "Implementation Status") is None
    assert _written(plan, "Implementation Date") is None
    assert any("audit activity reports" in s for s in plan.skipped)


def test_notes_come_from_the_capability_narrative_not_the_audit_log():
    plan = _plan(BLANK_ROW, _cap("Alpha", status="IMPLEMENTED", functions=[
        {"type": "PROVIDER", "name": "Capability Provider", "narrative": "Enforced."}]), None)
    # No audit activity, so no status or date -- but the narrative still lands.
    assert _written(plan, "Implementation Notes").new == "Enforced."
    assert _written(plan, "Implementation Status") is None


def test_fill_empty_will_not_replace_an_existing_note():
    row = dict(BLANK_ROW, **{"Implementation Notes": "Hand-written by the client."})
    plan = _plan(row, _cap("Alpha", functions=[
        {"type": "PROVIDER", "name": "P", "narrative": "From Paramify."}]))
    assert _written(plan, "Implementation Notes") is None


def test_sync_mode_does_replace_it():
    row = dict(BLANK_ROW, **{"Implementation Notes": "Hand-written by the client."})
    plan = _plan(row, _cap("Alpha", functions=[
        {"type": "PROVIDER", "name": "P", "narrative": "From Paramify."}]), mode=MODE_SYNC)
    assert _written(plan, "Implementation Notes").new == "From Paramify."


def test_no_matching_capability_writes_nothing():
    plan = _plan(BLANK_ROW, None)
    assert plan.updates == [] and plan.skipped


def test_status_that_would_invalidate_an_existing_pass_is_skipped():
    row = dict(BLANK_ROW, **{"Test Status": "Passed"})
    plan = _plan(row, _cap("Alpha"), _res("NOT_IMPLEMENTED"))
    assert _written(plan, "Implementation Status") is None
    assert any("permits Test Status" in s for s in plan.skipped)


def test_a_not_set_capability_still_publishes_its_narrative():
    # A narrative describes what the system DOES; the status column carries the
    # claim about whether it is in place. Suppressing the description because
    # the status is unset emptied the Notes column for 456 of 463 real rows.
    plan = _plan(BLANK_ROW, _cap("Alpha", status="NOT_SET", functions=[
        {"type": "PROVIDER", "name": "P", "narrative": "The platform enforces it."}]),
        _res("NOT_SET"))
    assert _written(plan, "Implementation Notes").new == "The platform enforces it."
    # ...while the STATUS is still refused, because NOT_SET has no Purview value
    assert _written(plan, "Implementation Status") is None


def test_skip_unassessed_notes_restores_the_cautious_behaviour():
    plan = plan_row(
        row_number=2, values=dict(BLANK_ROW),
        solcap=_cap("Alpha", status="NOT_SET", functions=[
            {"type": "PROVIDER", "name": "P", "narrative": "The platform enforces it."}]),
        resolution=None, tzinfo=UTC, skip_unassessed_notes=True,
    )
    assert _written(plan, "Implementation Notes") is None
    assert any("--skip-unassessed-notes" in s for s in plan.skipped)


def test_an_assessed_capability_is_unaffected_by_the_flag():
    plan = plan_row(
        row_number=2, values=dict(BLANK_ROW),
        solcap=_cap("Alpha", status="IMPLEMENTED", functions=[
            {"type": "PROVIDER", "name": "P", "narrative": "The platform enforces it."}]),
        resolution=None, tzinfo=UTC, skip_unassessed_notes=True,
    )
    assert _written(plan, "Implementation Notes").new == "The platform enforces it."


def test_partially_implemented_still_publishes_its_narrative():
    plan = _plan(BLANK_ROW, _cap("Alpha", status="PARTIALLY_IMPLEMENTED", functions=[
        {"type": "PROVIDER", "name": "P", "narrative": "Rolled out to two of three regions."}]))
    assert _written(plan, "Implementation Notes").new == "Rolled out to two of three regions."
    assert _written(plan, "Implementation Status") is None


# --- workbook safety ------------------------------------------------------- #

def test_workbook_refuses_to_write_a_read_only_column():
    import openpyxl
    import workbook as wb_mod
    from mapping import CellUpdate, RowPlan

    book = openpyxl.Workbook()
    sheet = book.active
    plan = RowPlan(row=2, action_id="a", action_name="Alpha")
    plan.updates.append(CellUpdate("Test Status", None, "Passed", True))
    with pytest.raises(wb_mod.WorkbookError, match="read-only"):
        wb_mod.apply(sheet, {"Test Status": 1}, [plan])


def test_workbook_rejects_a_short_export():
    import openpyxl
    import workbook as wb_mod

    book = openpyxl.Workbook()
    sheet = book.active
    for col, header in enumerate(["Action Id", "Improvement Action Name"], start=1):
        sheet.cell(row=1, column=col).value = header
    with pytest.raises(wb_mod.WorkbookError, match="missing required column"):
        wb_mod.header_index(sheet)


# --- end to end, offline --------------------------------------------------- #

def test_offline_run_against_the_real_export(tmp_path):
    import run as run_mod

    workbook = find_workbook()
    if workbook is None:
        pytest.skip("client export not present on this machine")
    code = run_mod.main([
        "--offline", str(TOOL / "fixtures"), "--workbook", str(workbook),
        "--out", str(tmp_path), "--no-run-dir", "--dry-run",
    ])
    assert code == 0
    report = json.loads((tmp_path / "run_report.json").read_text())
    assert report["dry_run"] is True
    assert report["totals"]["action_rows"] == 463
    assert report["totals"]["cells_by_column"] == {
        "Implementation Status": 2, "Implementation Date": 2, "Implementation Notes": 4,
    }
    # the orphan capability and the case-drift row are both surfaced
    assert report["join"]["unmatched_solution_capabilities"] == [
        "A Capability With No Purview Action"
    ]
    assert report["join"]["near_matches_detected"] == 1



# --- upload ---------------------------------------------------------------- #

def _report_stub(**over):
    base = {
        "generated_at": "2026-09-28T12:00:00Z",
        "mode": "fill-empty",
        "implementation_date_timezone": "UTC",
        "totals": {"cells_written": 7, "rows_with_writes": 3, "action_rows": 463,
                   "cells_by_column": {"Implementation Status": 2},
                   "cells_held_back": 1},
        "join": {"exact_matches": 6},
        "no_status_change_activity": {"rows": 4},
    }
    base.update(over)
    return base


def test_artifact_note_summarizes_the_run():
    import upload as upload_mod

    note = upload_mod.note_for(_report_stub())
    assert "7 cell(s) written" in note
    assert "audit log Activity feed" in note
    assert "no status-change" in note


def test_fixture_sourced_upload_is_labelled_as_a_test_artifact():
    # Uploading fixture-derived content into a real program without saying so
    # would seed it with something that looks like evidence.
    import upload as upload_mod

    note = upload_mod.note_for(_report_stub(fixture_sourced=True))
    assert note.startswith("*** TEST ARTIFACT")
    assert "NOT live Paramify data" in note


def test_upload_is_refused_with_dry_run(tmp_path):
    import run as run_mod

    workbook = find_workbook()
    if workbook is None:
        pytest.skip("client export not present on this machine")
    code = run_mod.main([
        "--offline", str(TOOL / "fixtures"), "--workbook", str(workbook),
        "--out", str(tmp_path), "--no-run-dir", "--dry-run", "--upload",
    ])
    assert code == 2


def test_upload_is_opt_in(tmp_path, monkeypatch):
    # A plain run must make no write call to Paramify at all.
    import run as run_mod
    import upload as upload_mod

    workbook = find_workbook()
    if workbook is None:
        pytest.skip("client export not present on this machine")
    called = []
    monkeypatch.setattr(upload_mod, "push", lambda **kw: called.append(kw))
    code = run_mod.main([
        "--offline", str(TOOL / "fixtures"), "--workbook", str(workbook),
        "--out", str(tmp_path), "--no-run-dir",
    ])
    assert code == 0 and called == []


# --- auth schemes ---------------------------------------------------------- #

def test_both_documented_auth_schemes_are_built_correctly(monkeypatch):
    from paramify_api import SCHEME_BEARER, SCHEME_LEGACY, ParamifyClient

    monkeypatch.setenv("PARAMIFY_API_TOKEN", "tok")
    client = ParamifyClient()
    assert client._headers(SCHEME_BEARER) == {"Authorization": "Bearer tok"}
    # apiKey_DEPRECATED: an API key in a header literally named "Bearer".
    assert client._headers(SCHEME_LEGACY) == {"Bearer": "tok"}


def test_a_401_on_the_current_scheme_retries_the_legacy_one(monkeypatch):
    from paramify_api import SCHEME_LEGACY, ParamifyClient

    monkeypatch.setenv("PARAMIFY_API_TOKEN", "tok")
    client = ParamifyClient()
    seen = []

    class Resp:
        def __init__(self, code):
            self.status_code, self.content = code, b'{"solutionCapabilities": []}'
            self.text = ""

        def json(self):
            return {"solutionCapabilities": []}

    def fake(method, url, headers=None, timeout=None, **kw):
        seen.append(headers)
        return Resp(401 if "Authorization" in headers else 200)

    monkeypatch.setattr(client._session, "request", fake)
    assert client.solution_capabilities() == []
    assert len(seen) == 2
    # and the working scheme is pinned, so later calls do not re-probe
    assert client.auth_scheme == SCHEME_LEGACY


def test_both_schemes_rejected_raises_an_actionable_error(monkeypatch):
    from paramify_api import ParamifyAuthError, ParamifyClient

    monkeypatch.setenv("PARAMIFY_API_TOKEN", "tok")
    client = ParamifyClient()

    class Resp:
        status_code, content, text = 401, b"{}", '{"message":"invalid"}'

    monkeypatch.setattr(client._session, "request",
                        lambda *a, **k: Resp())
    with pytest.raises(ParamifyAuthError, match="both auth schemes"):
        client.solution_capabilities()


def test_missing_token_names_the_fix(monkeypatch):
    from paramify_api import ParamifyClient, ParamifyError

    monkeypatch.delenv("PARAMIFY_API_TOKEN", raising=False)
    with pytest.raises(ParamifyError, match="export PARAMIFY_API_TOKEN"):
        ParamifyClient()


# --- full pipeline against a simulated API --------------------------------- #
# Proves everything downstream of authentication: fetch -> plan -> workbook ->
# evidence set -> multipart artifact upload. The only thing a real token adds is
# the 200 on the first call.

class _Resp:
    def __init__(self, payload, code=200):
        self._payload, self.status_code = payload, code
        self.content = b"{}" if payload is not None else b""
        self.text = ""

    def json(self):
        return self._payload


class _FakeSession:
    """Stands in for requests.Session, routing on the URL path."""

    def __init__(self):
        self.uploads = []
        self.created_evidence = []
        self._evidence = {}
        caps = json.loads((TOOL / "fixtures" / "solution_capabilities.json").read_text())
        events = json.loads((TOOL / "fixtures" / "audit_logs.json").read_text())
        self._caps, self._events = caps, events

    def request(self, method, url, headers=None, timeout=None, **kw):
        assert headers and ("Authorization" in headers or "Bearer" in headers)
        path = url.split("/api/v0", 1)[1]
        if method == "GET" and path.startswith("/solution-capabilities"):
            return _Resp(self._caps)
        if method == "GET" and path.startswith("/audit-logs"):
            return _Resp(self._events)
        if method == "GET" and path.startswith("/evidence"):
            ref = (kw.get("params") or {}).get("referenceId", [None])[0]
            hit = self._evidence.get(ref)
            return _Resp({"evidences": [hit] if hit else []})
        if method == "POST" and path == "/evidence":
            body = kw["json"]
            record = {"id": "ev-1", "referenceId": body["referenceId"],
                      "name": body["name"]}
            self._evidence[body["referenceId"]] = record
            self.created_evidence.append(record)
            return _Resp(record)
        if method == "POST" and "/artifacts/upload" in path:
            files = kw["files"]
            self.uploads.append({
                "evidence": path.split("/")[2],
                "filename": files["file"][0],
                "content_type": files["file"][2],
                "metadata": json.loads(files["artifact"][1]),
            })
            return _Resp({"id": f"art-{len(self.uploads)}"})
        raise AssertionError(f"unrouted call: {method} {path}")

    def close(self):
        pass


@pytest.fixture
def fake_api(monkeypatch):
    import paramify_api

    session = _FakeSession()
    monkeypatch.setenv("PARAMIFY_API_TOKEN", "test-token")
    monkeypatch.setattr(paramify_api.requests, "Session", lambda: session)
    return session


def _run(tmp_path, *extra):
    import run as run_mod

    workbook = find_workbook()
    if workbook is None:
        pytest.skip("client export not present on this machine")
    # --no-run-dir so these assert on a fixed path; the run-<ISO> directory has
    # its own tests below.
    return run_mod.main(
        ["--workbook", str(workbook), "--out", str(tmp_path), "--no-run-dir", *extra])


def test_full_pipeline_writes_the_workbook_and_uploads_both_artifacts(fake_api, tmp_path):
    # --min-match-rate 0: the fixtures deliberately cover only a few rows, and
    # this test is about the upload mechanics, not the join.
    assert _run(tmp_path, "--upload", "--min-match-rate", "0") == 0

    assert (tmp_path / "ExportActions.updated.xlsx").exists()
    assert len(fake_api.created_evidence) == 1
    assert fake_api.created_evidence[0]["referenceId"] == "PURVIEW-ACTION-UPDATE"

    names = [u["filename"] for u in fake_api.uploads]
    assert names == ["ExportActions.updated.xlsx", "run_report.json"]
    xlsx = fake_api.uploads[0]
    assert xlsx["content_type"].endswith("spreadsheetml.sheet")
    assert "cell(s) written" in xlsx["metadata"]["note"]
    assert xlsx["metadata"]["effectiveDate"].endswith("Z")


def test_rerunning_reuses_the_evidence_set_instead_of_duplicating_it(fake_api, tmp_path):
    assert _run(tmp_path, "--upload", "--min-match-rate", "0") == 0
    assert _run(tmp_path, "--upload", "--min-match-rate", "0") == 0
    assert len(fake_api.created_evidence) == 1      # created once
    assert len(fake_api.uploads) == 4               # artifacts append, by design


def test_live_run_derives_status_and_date_from_the_audit_activity(fake_api, tmp_path):
    assert _run(tmp_path) == 0
    report = json.loads((tmp_path / "run_report.json").read_text())
    assert report["sources"]["implementation_status"] == "audit_log_activity"
    assert report["totals"]["cells_by_column"] == {
        "Implementation Status": 2, "Implementation Date": 2, "Implementation Notes": 4,
    }
    assert report.get("upload") is None             # no upload without the flag


def test_a_filename_needing_encoding_is_url_encoded(fake_api, tmp_path):
    import shutil

    source = find_workbook()
    if source is None:
        pytest.skip("client export not present on this machine")
    awkward = tmp_path / "Export Actions (CMMC L2) #1.xlsx"
    shutil.copy(source, awkward)

    import run as run_mod
    assert run_mod.main(["--workbook", str(awkward), "--out", str(tmp_path / "o"),
                         "--no-run-dir", "--upload", "--min-match-rate", "0"]) == 0
    assert fake_api.uploads[0]["filename"] == (
        "Export%20Actions%20%28CMMC%20L2%29%20%231.updated.xlsx"
    )


# --- base URL selection ---------------------------------------------------- #

def test_base_url_is_resolved_at_call_time_not_import_time(monkeypatch):
    # A module-level constant would freeze before .env is loaded, so a base URL
    # set there would be ignored -- failing as a 401 that looks like a bad token.
    import config

    monkeypatch.delenv("PARAMIFY_API_BASE_URL", raising=False)
    assert config.base_url() == config.DEFAULT_BASE_URL
    monkeypatch.setenv("PARAMIFY_API_BASE_URL", "https://stage.paramify.com/api/v0/")
    assert config.base_url() == "https://stage.paramify.com/api/v0"


def test_client_picks_up_the_staging_base_url(monkeypatch):
    from paramify_api import ParamifyClient

    monkeypatch.setenv("PARAMIFY_API_TOKEN", "tok")
    monkeypatch.setenv("PARAMIFY_API_BASE_URL", "https://stage.paramify.com/api/v0")
    assert ParamifyClient()._base_url == "https://stage.paramify.com/api/v0"


def test_base_url_flag_wins_over_the_environment(monkeypatch, tmp_path):
    import run as run_mod

    monkeypatch.setenv("PARAMIFY_API_TOKEN", "tok")
    monkeypatch.setenv("PARAMIFY_API_BASE_URL", "https://app.paramify.com/api/v0")
    seen = {}

    import paramify_api

    class Cap:
        def __init__(self, *a, **k):
            seen["url"] = paramify_api.ParamifyClient()._base_url
            raise SystemExit(0)

    monkeypatch.setattr(run_mod, "ParamifyClient", Cap)
    with pytest.raises(SystemExit):
        run_mod.main(["--check-auth", "--base-url", "https://stage.paramify.com/api/v0"])
    assert seen["url"] == "https://stage.paramify.com/api/v0"


# --- wrong-workspace guard ------------------------------------------------- #

def test_a_partial_mapping_uploads_without_complaint(fake_api, tmp_path):
    """The counts legitimately differ.

    A Paramify program holds capabilities for the improvement actions it covers;
    the rest of a Purview assessment simply has none. The fixtures match 6 of
    463 rows, which is exactly that shape, and must not be treated as a fault.
    """
    assert _run(tmp_path, "--upload") == 0
    assert len(fake_api.uploads) == 2


def test_the_match_rate_guard_applies_only_when_asked_for(fake_api, tmp_path, caplog):
    code = _run(tmp_path, "--upload", "--min-match-rate", "0.5")
    assert code == 3
    assert fake_api.uploads == []
    assert "below the --min-match-rate you set" in caplog.text
    # still written locally for diagnosis
    assert (tmp_path / "run_report.json").exists()


def test_zero_matches_is_called_out_as_an_actual_problem(fake_api, tmp_path, monkeypatch):
    # Some matches is normal; none means the token is on the wrong workspace.
    import report as report_mod
    from matching import MatchResult

    built = report_mod.build(
        plans=[], matches=MatchResult(), mode="fill-empty", timezone_name="UTC",
        workbook_in="in.xlsx", workbook_out=None, audit_status={"state": "x"},
        capability_count=0, dry_run=True,
    )
    assert "NOTHING matched by name" in report_mod.summarize(built)


def test_the_guard_can_be_disabled(fake_api, tmp_path):
    assert _run(tmp_path, "--upload", "--min-match-rate", "0") == 0
    assert len(fake_api.uploads) == 2


def test_unmatched_counts_are_reported_as_facts_not_failures(fake_api, tmp_path):
    import json as _json

    assert _run(tmp_path) == 0
    report = _json.loads((tmp_path / "run_report.json").read_text())
    # both directions are recorded, and neither stops the run
    assert report["join"]["unmatched_action_rows"]
    assert report["join"]["unmatched_solution_capabilities"]
    assert report["join"]["exact_matches"] > 0


def test_value_strategy_requires_both_sides_to_be_status_tokens():
    # Observed shape: {"old": "NOT_SET", "new": "IMPLEMENTED"} -- accepted.
    got = audit.index_by_capability([
        _event("c1", "2026-09-28T10:00:00Z", [{"old": "NOT_SET", "new": "IMPLEMENTED"}])
    ])
    assert got["c1"].status_from_audit == "IMPLEMENTED"


def test_a_field_edited_to_status_like_text_is_not_a_status_change():
    # Without a field name, this is the strategy's one false positive; requiring
    # `old` to qualify too removes it.
    assert audit.index_by_capability([
        _event("c1", "2026-09-28T10:00:00Z",
               [{"old": "Some narrative prose", "new": "IMPLEMENTED"}])
    ]) == {}


def test_a_first_transition_from_empty_still_resolves():
    got = audit.index_by_capability([
        _event("c1", "2026-09-28T10:00:00Z", [{"old": "", "new": "IMPLEMENTED"}])
    ])
    assert got["c1"].status_from_audit == "IMPLEMENTED"


def test_upload_is_refused_when_the_run_changed_nothing(fake_api, tmp_path, caplog, monkeypatch):
    # An unchanged workbook in an evidence set is noise, not evidence.
    import planner

    monkeypatch.setattr(planner, "_should_write", lambda e, n, m: (False, "test: held"))
    code = _run(tmp_path, "--upload", "--min-match-rate", "0")
    assert code == 4
    assert fake_api.uploads == []
    assert "changed no cells" in caplog.text


def test_a_no_op_upload_can_be_forced(fake_api, tmp_path, monkeypatch):
    import planner

    monkeypatch.setattr(planner, "_should_write", lambda e, n, m: (False, "test: held"))
    assert _run(tmp_path, "--upload", "--min-match-rate", "0", "--allow-no-changes") == 0
    assert len(fake_api.uploads) == 2


# --- plain-text normalization ---------------------------------------------- #

def test_the_mention_at_is_dropped_but_the_name_is_kept():
    from mapping import plain_text

    assert plain_text(
        "@Microsoft 365 Government Community Cloud-High captures the execution of "
        "privileged functions."
    ) == (
        "Microsoft 365 Government Community Cloud-High captures the execution of "
        "privileged functions."
    )


def test_an_email_address_keeps_its_at_sign():
    from mapping import plain_text

    assert plain_text("Notifications go to soc@example.com daily.") == (
        "Notifications go to soc@example.com daily."
    )


def test_whitespace_is_tidied_but_paragraphs_survive():
    from mapping import plain_text

    assert plain_text("  One   two \r\n\n\n\n  three  ") == "One two\n\nthree"


@pytest.mark.parametrize("value", ["", None, "   \n  "])
def test_empty_narratives_normalize_to_empty(value):
    from mapping import plain_text

    assert plain_text(value) == ""


def test_a_rich_text_document_is_refused_not_written():
    # Paramify stores narratives twice; only the flattened form belongs in a cell.
    from mapping import looks_like_rich_text

    doc = '[{"type":"p","children":[{"text":""},{"type":"mention","value":"X"}]}]'
    assert looks_like_rich_text(doc)
    assert render_notes([{"type": "PROVIDER", "name": "P", "narrative": doc}]) == (None, [])


def test_notes_rendered_for_a_cell_carry_no_at_mention():
    text, _ = render_notes([{
        "type": "PROVIDER", "name": "Capability Provider",
        "narrative": "@Microsoft 365 Government Community Cloud-High labels ePHI.",
    }])
    assert text.startswith("Microsoft 365")
    assert "@" not in text


# --- Test Date conflict resolution ----------------------------------------- #

ROW_420 = {
    "Action Id": "a", "Improvement Action Name": "Alert personnel of information spillage",
    "Implementation Status": "Implemented", "Implementation Date": "6/19/2026 15:12:29",
    "Implementation Notes": "SSP and IR plan cover information spillage.",
    "Test Status": "Passed", "Test Date": "6/19/2026 15:12:29",
}
LATER = audit.DateResolution(
    datetime(2026, 9, 28, 21, 32, 58, tzinfo=timezone.utc), "IMPLEMENTED", "value")


def _conflict_plan(choice, tz="UTC"):
    from planner import MODE_SYNC

    return plan_row(row_number=420, values=dict(ROW_420),
                    solcap=_cap("Alert personnel of information spillage"),
                    resolution=LATER, tzinfo=ZoneInfo(tz), mode=MODE_SYNC,
                    on_date_conflict=choice)


def test_skip_is_the_default_and_writes_no_date():
    from planner import CONFLICT_SKIP

    plan = _conflict_plan(CONFLICT_SKIP)
    assert _written(plan, "Implementation Date") is None
    assert any("--on-date-conflict" in s for s in plan.skipped)


def test_advance_test_moves_the_test_date_with_the_implementation_date():
    from planner import CONFLICT_ADVANCE_TEST

    plan = _conflict_plan(CONFLICT_ADVANCE_TEST)
    assert _written(plan, "Implementation Date").new == "9/28/2026 21:32:58"
    assert _written(plan, "Test Date").new == "9/28/2026 21:32:58"
    # the recorded result survives, but the row now claims a test that day
    assert _written(plan, "Test Status") is None
    assert any(a["kind"] == "test_date_advanced"
               and "no longer reflects when the test was performed" in a["message"]
               for a in plan.advisories)


def test_clear_test_discards_the_stale_result_and_says_so():
    from planner import CONFLICT_CLEAR_TEST

    plan = _conflict_plan(CONFLICT_CLEAR_TEST)
    assert _written(plan, "Implementation Date").new == "9/28/2026 21:32:58"
    assert _written(plan, "Test Date").new is None
    assert _written(plan, "Test Status").new is None
    assert any(a["kind"] == "test_result_cleared"
               and "needs retesting" in a["message"] for a in plan.advisories)


def test_the_test_columns_are_untouched_when_the_date_is_not_written():
    # fill-empty leaves the occupied date alone, so clearing the test beside it
    # would make the row incoherent.
    from planner import CONFLICT_CLEAR_TEST, MODE_FILL_EMPTY

    plan = plan_row(row_number=420, values=dict(ROW_420),
                    solcap=_cap("Alert personnel of information spillage"),
                    resolution=LATER, tzinfo=ZoneInfo("UTC"), mode=MODE_FILL_EMPTY,
                    on_date_conflict=CONFLICT_CLEAR_TEST)
    assert _written(plan, "Implementation Date") is None
    assert _written(plan, "Test Date") is None
    assert _written(plan, "Test Status") is None


def test_timezone_renders_the_date_as_paramify_displays_it():
    from planner import CONFLICT_ADVANCE_TEST

    # The Activity feed shows 3:32 PM Mountain for this event.
    plan = _conflict_plan(CONFLICT_ADVANCE_TEST, tz="America/Denver")
    assert _written(plan, "Implementation Date").new == "9/28/2026 15:32:58"


def test_test_columns_stay_unwritable_without_the_flag():
    import openpyxl
    import workbook as wb_mod
    from mapping import CellUpdate, RowPlan

    book = openpyxl.Workbook()
    plan = RowPlan(row=2, action_id="a", action_name="x")
    plan.updates.append(CellUpdate("Test Date", None, "9/28/2026 15:32:58", True))
    with pytest.raises(wb_mod.WorkbookError, match="read-only"):
        wb_mod.apply(book.active, {"Test Date": 1}, [plan])


def test_advisories_are_bucketed_by_kind_not_lumped_together():
    # Regression: the report counted EVERY advisory as a blank-Test-Status
    # warning, so resolving a date conflict reported 6 rules-ambiguity rows
    # that did not exist. A miscounted caveat is worse than no caveat.
    import report as report_mod
    from matching import MatchResult
    from planner import CONFLICT_ADVANCE_TEST, MODE_SYNC

    plan = plan_row(row_number=420, values=dict(ROW_420),
                    solcap=_cap("Alert personnel of information spillage"),
                    resolution=LATER, tzinfo=ZoneInfo("UTC"), mode=MODE_SYNC,
                    on_date_conflict=CONFLICT_ADVANCE_TEST)
    built = report_mod.build(
        plans=[plan], matches=MatchResult(), mode="sync", timezone_name="UTC",
        workbook_in="in.xlsx", workbook_out="out.xlsx", audit_status={"state": "x"},
        capability_count=1, dry_run=False,
    )
    assert set(built["advisories"]) == {"test_date_advanced"}
    assert built["advisories"]["test_date_advanced"]["rows"] == 1
    # the row's Test Status was 'Passed', so this bucket must be absent entirely
    assert "blank_test_status" not in built["advisories"]


# --- timezone: only the DATE has to be right ------------------------------- #

def test_utc_would_put_a_us_evening_change_on_the_wrong_date():
    # 2026-09-29T02:00Z is still 28 September everywhere in the US.
    from planner import CONFLICT_ADVANCE_TEST, MODE_SYNC

    evening = audit.DateResolution(
        datetime(2026, 9, 29, 2, 0, 0, tzinfo=timezone.utc), "IMPLEMENTED", "value")

    def date_in(tz):
        plan = plan_row(row_number=1, values=dict(ROW_420),
                        solcap=_cap("x"), resolution=evening, tzinfo=ZoneInfo(tz),
                        mode=MODE_SYNC, on_date_conflict=CONFLICT_ADVANCE_TEST)
        return _written(plan, "Implementation Date").new.split(" ")[0]

    assert date_in("UTC") == "9/29/2026"                 # wrong for a US operator
    for zone in ("America/New_York", "America/Chicago", "America/Denver",
                 "America/Phoenix", "America/Los_Angeles", "America/Anchorage",
                 "Pacific/Honolulu"):
        assert date_in(zone) == "9/28/2026", zone


@pytest.mark.parametrize("zone", [
    "America/New_York", "America/Chicago", "America/Denver", "America/Phoenix",
    "America/Los_Angeles", "America/Anchorage", "Pacific/Honolulu",
])
def test_every_us_zone_resolves(zone):
    import config

    tzinfo, name, how = config.resolve_timezone(zone)
    assert name == zone and how == "--timezone"


def test_timezone_defaults_to_the_operators_own_zone(monkeypatch):
    import config

    monkeypatch.delenv("TZ", raising=False)
    _, name, how = config.resolve_timezone(None)
    # Not UTC-by-default: that is what put dates on the wrong day.
    assert how in {"system zone (/etc/localtime)", "system UTC offset"}
    assert name


def test_the_tz_environment_variable_is_honoured(monkeypatch):
    import config

    monkeypatch.setenv("TZ", "America/Chicago")
    _, name, how = config.resolve_timezone(None)
    assert name == "America/Chicago"
    assert how == "TZ environment variable"


def test_an_explicit_zone_beats_the_environment(monkeypatch):
    import config

    monkeypatch.setenv("TZ", "America/Chicago")
    _, name, _ = config.resolve_timezone("America/New_York")
    assert name == "America/New_York"


def test_local_forces_detection_even_with_tz_set(monkeypatch):
    import config

    monkeypatch.setenv("TZ", "America/Chicago")
    _, _, how = config.resolve_timezone("local")
    assert how != "TZ environment variable"


def test_an_unknown_zone_is_refused_with_the_us_list():
    import config

    with pytest.raises(config.ConfigError, match="America/Denver"):
        config.resolve_timezone("Mars/Olympus")


# --- notes policy: keep both --------------------------------------------- #

CLIENT_NOTE = "Created CUI sensitivity Label and enforced policy"
NARRATIVE = "Microsoft 365 GCC High labels ePHI automatically via auto-labeling policy."


def _notes_plan(existing, policy, mode=MODE_FILL_EMPTY):
    from planner import NOTES_APPEND, NOTES_FOLLOW_MODE  # noqa: F401

    row = dict(BLANK_ROW, **{"Implementation Notes": existing})
    return plan_row(
        row_number=2, values=row,
        solcap=_cap("Alpha", functions=[
            {"type": "PROVIDER", "name": "Capability Provider", "narrative": NARRATIVE}]),
        resolution=None, tzinfo=ZoneInfo("UTC"), mode=mode, notes_policy=policy,
    )


def test_append_keeps_the_clients_note_and_adds_the_narrative():
    from planner import NOTES_APPEND

    written = _written(_notes_plan(CLIENT_NOTE, NOTES_APPEND), "Implementation Notes")
    assert written.new == CLIENT_NOTE + "\n\n" + NARRATIVE
    assert CLIENT_NOTE in written.new and NARRATIVE in written.new


def test_append_is_idempotent_across_reruns():
    # Without this, the fifth run leaves the cell unreadable.
    from planner import NOTES_APPEND

    once = _written(_notes_plan(CLIENT_NOTE, NOTES_APPEND), "Implementation Notes").new
    twice = _notes_plan(once, NOTES_APPEND)
    assert _written(twice, "Implementation Notes") is None
    assert any("already present" in (u.reason or "") for u in twice.updates)


def test_append_writes_the_narrative_alone_into_a_blank_cell():
    from planner import NOTES_APPEND

    assert _written(_notes_plan(None, NOTES_APPEND), "Implementation Notes").new == NARRATIVE


def test_follow_mode_is_unchanged_by_the_new_policy():
    from planner import MODE_SYNC, NOTES_FOLLOW_MODE

    assert _written(_notes_plan(CLIENT_NOTE, NOTES_FOLLOW_MODE), "Implementation Notes") is None
    replaced = _written(
        _notes_plan(CLIENT_NOTE, NOTES_FOLLOW_MODE, mode=MODE_SYNC), "Implementation Notes")
    assert replaced.new == NARRATIVE          # the client's note is gone
    assert CLIENT_NOTE not in replaced.new


def test_append_survives_a_sync_run_without_destroying_anything():
    # The combination a production run would actually use.
    from planner import MODE_SYNC, NOTES_APPEND

    written = _written(
        _notes_plan(CLIENT_NOTE, NOTES_APPEND, mode=MODE_SYNC), "Implementation Notes")
    assert written.new.startswith(CLIENT_NOTE)


# --- cleanup: deletes are gated and identified by what the run wrote ------- #

class _CleanupSession:
    def __init__(self, artifacts):
        self._artifacts = list(artifacts)
        self.deleted_artifacts, self.deleted_sets = [], []

    def request(self, method, url, headers=None, timeout=None, **kw):
        path = url.split("/api/v0", 1)[1]
        if method == "GET" and path.startswith("/evidence?") or path == "/evidence":
            ref = (kw.get("params") or {}).get("referenceId", [None])[0]
            hit = {"id": "ev-known", "referenceId": ref} if ref == "PURVIEW-ACTION-UPDATE" else None
            return _Resp({"evidences": [hit] if hit else []})
        if method == "GET" and path.endswith("/artifacts"):
            return _Resp({"artifacts": self._artifacts})
        if method == "DELETE" and "/artifacts/" in path:
            self.deleted_artifacts.append(path.rsplit("/", 1)[1])
            return _Resp(None)
        if method == "DELETE":
            self.deleted_sets.append(path.rsplit("/", 1)[1])
            return _Resp(None)
        raise AssertionError(f"unrouted: {method} {path}")

    def close(self):
        pass


ARTIFACTS = [
    {"id": "a1", "title": "ExportActions.updated.xlsx", "effectiveDate": "2026-09-28",
     "note": "Generated by tools/purview_action_update. 0 cell(s) written across 0 rows."},
    {"id": "a2", "title": "run_report.json", "effectiveDate": "2026-09-28",
     "note": "Run report ... 0 cell(s) written ..."},
    {"id": "a3", "title": "ExportActions.updated.xlsx", "effectiveDate": "2026-09-28",
     "note": "Generated by tools/purview_action_update. 19 cell(s) written across 7 rows."},
]


@pytest.fixture
def cleanup_api(monkeypatch):
    import paramify_api

    session = _CleanupSession(ARTIFACTS)
    monkeypatch.setenv("PARAMIFY_API_TOKEN", "tok")
    monkeypatch.setattr(paramify_api.requests, "Session", lambda: session)
    return session


def test_pruning_reports_before_it_deletes(cleanup_api):
    import upload as upload_mod

    result = upload_mod.prune_no_op_artifacts("ev-1", confirm=False)
    assert result["would_delete"] == ["a1", "a2"]
    assert result["deleted"] == []
    assert cleanup_api.deleted_artifacts == []


def test_pruning_deletes_only_the_no_op_artifacts(cleanup_api):
    import upload as upload_mod

    result = upload_mod.prune_no_op_artifacts("ev-1", confirm=True)
    assert result["deleted"] == ["a1", "a2"]
    # the artifact from the run that actually changed 19 cells survives
    assert "a3" not in cleanup_api.deleted_artifacts


def test_no_ops_are_identified_by_the_note_not_by_position(cleanup_api):
    # A human-uploaded artifact has no such note and must never be selected.
    import upload as upload_mod

    manual = {"id": "a4", "title": "evidence.pdf", "effectiveDate": "2026-09-01",
              "note": "Uploaded by the assessor."}
    cleanup_api._artifacts = [manual] + ARTIFACTS
    assert "a4" not in upload_mod.prune_no_op_artifacts("ev-1", confirm=False)["would_delete"]


def test_deleting_an_evidence_set_needs_confirm(cleanup_api):
    import upload as upload_mod

    assert upload_mod.delete_evidence_set("ev-1", confirm=False)["deleted"] is None
    assert cleanup_api.deleted_sets == []
    assert upload_mod.delete_evidence_set("ev-1", confirm=True)["deleted"] == "ev-1"
    assert cleanup_api.deleted_sets == ["ev-1"]


def test_describe_flags_which_artifacts_are_no_ops():
    import upload as upload_mod

    lines = {line.split()[0]: line for line in upload_mod.describe(ARTIFACTS).splitlines()}
    assert "<- no-op" in lines["a1"]
    assert "<- no-op" in lines["a2"]
    assert "<- no-op" not in lines["a3"]   # the run that wrote 19 cells


# --- post-write verification ----------------------------------------------- #
# A verifier that never fails is decoration. Each case below corrupts the
# produced workbook in one specific way and asserts it is caught.

VERIFY_HEADERS = [
    "Action Id", "Improvement Action Name", "Applicable To", "Service Scope", "Group",
    "Implementation Status", "Implementation Date", "Implementation Notes",
    "Test Status", "Test Date", "Test Notes", "Other Notes", "Documents",
    "Assigned To", "Testing Type",
]
VERIFY_ROW = [
    "id-1", "Alpha", "Microsoft 365", "N/A", "Default Group",
    "Implemented", "6/19/2026 15:12:29", "A note",
    "Passed", "6/19/2026 15:12:29", None, None, None, None, "Manual",
]


def _book(path, rows=None, extra_sheet=True):
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Action Update"
    sheet.append(VERIFY_HEADERS)
    for row in (rows or [list(VERIFY_ROW)]):
        sheet.append(row)
    if extra_sheet:
        ref = book.create_sheet("How To Update Actions")
        ref.append(["rules", "live", "here"])
    book.save(path)
    return path


WRITABLE3 = ("Implementation Status", "Implementation Date", "Implementation Notes")
WRITABLE5 = WRITABLE3 + ("Test Date", "Test Status")


def test_an_untouched_copy_verifies_clean(tmp_path):
    import verify as verify_mod

    a, b = _book(tmp_path / "a.xlsx"), _book(tmp_path / "b.xlsx")
    assert verify_mod.verify(a, b, writable=WRITABLE3) == []


def test_a_legitimate_edit_verifies_clean(tmp_path):
    import verify as verify_mod

    row = list(VERIFY_ROW)
    row[7] = "A replacement note"          # Implementation Notes, a writable column
    a, b = _book(tmp_path / "a.xlsx"), _book(tmp_path / "b.xlsx", [row])
    assert verify_mod.verify(a, b, writable=WRITABLE3) == []


def _checks(problems):
    return {p["check"] for p in problems}


def test_it_catches_a_write_to_a_read_only_column(tmp_path):
    import verify as verify_mod

    row = list(VERIFY_ROW)
    row[12] = "doc::http://example.test"   # Documents — never writable
    a, b = _book(tmp_path / "a.xlsx"), _book(tmp_path / "b.xlsx", [row])
    assert "read-only-column" in _checks(verify_mod.verify(a, b, writable=WRITABLE3))


def test_it_catches_a_test_date_before_the_implementation_date(tmp_path):
    import verify as verify_mod

    row = list(VERIFY_ROW)
    row[6] = "9/28/2026 15:32:58"          # implementation moved past the June test
    a, b = _book(tmp_path / "a.xlsx"), _book(tmp_path / "b.xlsx", [row])
    assert "test-date-order" in _checks(verify_mod.verify(a, b, writable=WRITABLE3))


def test_it_catches_an_impossible_status_combination(tmp_path):
    import verify as verify_mod

    row = list(VERIFY_ROW)
    row[5] = "NotImplemented"              # NotImplemented permits only Test Status None
    a, b = _book(tmp_path / "a.xlsx"), _book(tmp_path / "b.xlsx", [row])
    assert "status-combination" in _checks(verify_mod.verify(a, b, writable=WRITABLE3))


def test_it_catches_a_status_outside_purviews_vocabulary(tmp_path):
    import verify as verify_mod

    row = list(VERIFY_ROW)
    row[5], row[8], row[9] = "IMPLEMENTED", None, None   # the Paramify spelling
    a, b = _book(tmp_path / "a.xlsx"), _book(tmp_path / "b.xlsx", [row])
    assert "status-vocabulary" in _checks(verify_mod.verify(a, b, writable=WRITABLE3))


def test_it_catches_a_mangled_date_format(tmp_path):
    import verify as verify_mod

    row = list(VERIFY_ROW)
    row[6], row[8], row[9] = "2026-09-28T15:32:58Z", None, None
    a, b = _book(tmp_path / "a.xlsx"), _book(tmp_path / "b.xlsx", [row])
    assert "date-format" in _checks(verify_mod.verify(a, b, writable=WRITABLE3))


def test_it_catches_a_clobbered_reference_sheet(tmp_path):
    import openpyxl
    import verify as verify_mod

    a, b = _book(tmp_path / "a.xlsx"), _book(tmp_path / "b.xlsx")
    book = openpyxl.load_workbook(b)
    book["How To Update Actions"]["A1"] = "clobbered"
    book.save(b)
    assert "untouched-sheet" in _checks(verify_mod.verify(a, b, writable=WRITABLE3))


def test_it_catches_a_changed_action_id(tmp_path):
    import verify as verify_mod

    row = list(VERIFY_ROW)
    row[0] = "id-2"
    a, b = _book(tmp_path / "a.xlsx"), _book(tmp_path / "b.xlsx", [row])
    assert "identity-changed" in _checks(verify_mod.verify(a, b, writable=WRITABLE3))


def test_it_catches_a_dropped_row(tmp_path):
    import verify as verify_mod

    two = [list(VERIFY_ROW), list(VERIFY_ROW)]
    two[1][0] = "id-2"
    a = _book(tmp_path / "a.xlsx", two)
    b = _book(tmp_path / "b.xlsx", [list(VERIFY_ROW)])
    assert "dimensions" in _checks(verify_mod.verify(a, b, writable=WRITABLE3))


def test_advance_test_output_verifies_clean_under_the_wider_writable_set(tmp_path):
    import verify as verify_mod

    row = list(VERIFY_ROW)
    row[6] = row[9] = "9/28/2026 15:32:58"   # both dates moved together
    a, b = _book(tmp_path / "a.xlsx"), _book(tmp_path / "b.xlsx", [row])
    assert verify_mod.verify(a, b, writable=WRITABLE5) == []
    # but the same file is a violation when Test Date was not meant to be writable
    assert "read-only-column" in _checks(verify_mod.verify(a, b, writable=WRITABLE3))


def test_the_changes_preview_shows_paragraph_breaks():
    # Collapsing "\n\n" to a space made an appended note look like two
    # sentences jammed together, when the cell was correct.
    import report as report_mod

    assert "⏎⏎" in report_mod._short("Client note.\n\nParamify narrative.", limit=200)
    assert report_mod._short(None) == "(blank)"


def test_replace_overwrites_the_clients_note_entirely():
    from planner import NOTES_REPLACE

    written = _written(_notes_plan(CLIENT_NOTE, NOTES_REPLACE), "Implementation Notes")
    assert written.new == NARRATIVE
    assert CLIENT_NOTE not in written.new


def test_replace_ignores_mode_so_fill_empty_still_protects_status_and_dates():
    from planner import MODE_FILL_EMPTY, NOTES_REPLACE

    written = _written(
        _notes_plan(CLIENT_NOTE, NOTES_REPLACE, mode=MODE_FILL_EMPTY),
        "Implementation Notes")
    assert written.new == NARRATIVE


def test_replace_is_a_no_op_when_the_cell_already_holds_the_narrative():
    from planner import NOTES_REPLACE

    plan = _notes_plan(NARRATIVE, NOTES_REPLACE)
    assert _written(plan, "Implementation Notes") is None


# --- run directories and input provenance ---------------------------------- #
# On a recurring cadence the workbook must come FRESH from Purview each cycle.
# Re-using the last run's output silently discards anything changed in Purview
# since — a new action, an edited note, a recorded test.

def test_each_run_gets_its_own_timestamped_directory(fake_api, tmp_path):
    import run as run_mod

    workbook = find_workbook()
    if workbook is None:
        pytest.skip("client export not present on this machine")
    assert run_mod.main(["--workbook", str(workbook), "--out", str(tmp_path)]) == 0
    dirs = [d for d in tmp_path.iterdir() if d.is_dir() and d.name.startswith("run-")]
    assert len(dirs) == 1
    assert (dirs[0] / "run_report.json").exists()
    assert (dirs[0] / f"{workbook.stem}.updated.xlsx").exists()


def test_a_second_run_does_not_overwrite_the_first(fake_api, tmp_path, monkeypatch):
    import run as run_mod

    workbook = find_workbook()
    if workbook is None:
        pytest.skip("client export not present on this machine")

    class _Clock:
        """Stands in for `datetime` so two runs land in distinct directories."""

        def __init__(self, moments):
            self._moments = iter(moments)

        def now(self, tz=None):
            return next(self._moments)

    monkeypatch.setattr(run_mod, "datetime", _Clock([
        datetime(2026, 9, 29, 10, 0, 0, tzinfo=timezone.utc),
        datetime(2026, 10, 2, 10, 0, 0, tzinfo=timezone.utc),
    ]))
    for _ in range(2):
        assert run_mod.main(["--workbook", str(workbook), "--out", str(tmp_path)]) == 0

    names = sorted(d.name for d in tmp_path.iterdir() if d.is_dir())
    assert names == ["run-2026-09-29T10-00-00Z", "run-2026-10-02T10-00-00Z"]
    # both cycles' uploaded workbook and report are still on disk
    assert all((tmp_path / n / "run_report.json").exists() for n in names)


def test_the_produced_workbook_is_stamped_as_a_tool_output(fake_api, tmp_path):
    import run as run_mod
    import workbook as wb_mod

    workbook = find_workbook()
    if workbook is None:
        pytest.skip("client export not present on this machine")
    run_mod.main(["--workbook", str(workbook), "--out", str(tmp_path), "--no-run-dir"])
    stamp = wb_mod.provenance_of(tmp_path / f"{workbook.stem}.updated.xlsx")
    assert stamp and stamp[wb_mod.PROP_TOOL] == "purview_action_update"


def test_a_fresh_purview_export_carries_no_stamp():
    import workbook as wb_mod

    workbook = find_workbook()
    if workbook is None:
        pytest.skip("client export not present on this machine")
    assert wb_mod.provenance_of(workbook) is None


def test_feeding_a_previous_output_back_in_is_refused(fake_api, tmp_path, caplog):
    import run as run_mod

    workbook = find_workbook()
    if workbook is None:
        pytest.skip("client export not present on this machine")
    run_mod.main(["--workbook", str(workbook), "--out", str(tmp_path), "--no-run-dir"])
    produced = tmp_path / f"{workbook.stem}.updated.xlsx"

    code = run_mod.main(
        ["--workbook", str(produced), "--out", str(tmp_path / "second"), "--no-run-dir"])
    assert code == 6
    assert "FRESH export" in caplog.text
    assert not (tmp_path / "second").exists()


def test_the_chain_guard_can_be_overridden(fake_api, tmp_path):
    import run as run_mod

    workbook = find_workbook()
    if workbook is None:
        pytest.skip("client export not present on this machine")
    run_mod.main(["--workbook", str(workbook), "--out", str(tmp_path), "--no-run-dir"])
    produced = tmp_path / f"{workbook.stem}.updated.xlsx"
    assert run_mod.main(["--workbook", str(produced), "--out", str(tmp_path / "s"),
                         "--no-run-dir", "--allow-chained-input"]) == 0


def test_stamping_changes_no_cell(tmp_path):
    # The stamp lives in package metadata, never in a sheet Purview reads.
    import openpyxl
    import workbook as wb_mod

    workbook = find_workbook()
    if workbook is None:
        pytest.skip("client export not present on this machine")
    book = openpyxl.load_workbook(workbook)
    wb_mod.stamp_provenance(book, run_id="run-x", source="src.xlsx")
    out = tmp_path / "stamped.xlsx"
    book.save(out)
    a = openpyxl.load_workbook(workbook, data_only=True)
    b = openpyxl.load_workbook(out, data_only=True)
    assert a.sheetnames == b.sheetnames
    assert all(a[s].cell(r, c).value == b[s].cell(r, c).value
               for s in a.sheetnames
               for r in range(1, a[s].max_row + 1)
               for c in range(1, a[s].max_column + 1))


# --- the working area can live anywhere --------------------------------- #

def _seed_inbox(base):
    src = find_workbook()
    if src is None:
        pytest.skip("client export not present on this machine")
    import shutil
    inbox = base / "inbox"
    inbox.mkdir(parents=True, exist_ok=True)
    shutil.copy(src, inbox / "ExportActions.xlsx")
    return src


def test_a_working_area_anywhere_drives_both_paths(fake_api, tmp_path, monkeypatch):
    """One setting, no path arguments. The folder may be anywhere."""
    import run as run_mod

    base = tmp_path / "somewhere" / "deeply" / "nested"
    _seed_inbox(base)
    monkeypatch.setenv("PURVIEW_WRITEBACK_DIR", str(base))
    assert run_mod.main([]) == 0

    runs = [d for d in (base / "runs").iterdir() if d.name.startswith("run-")]
    assert len(runs) == 1
    assert (runs[0] / "ExportActions.updated.xlsx").exists()
    assert (runs[0] / "run_report.json").exists()


def test_the_flag_beats_the_environment(fake_api, tmp_path, monkeypatch):
    import run as run_mod

    chosen, ignored = tmp_path / "chosen", tmp_path / "ignored"
    _seed_inbox(chosen)
    monkeypatch.setenv("PURVIEW_WRITEBACK_DIR", str(ignored))
    assert run_mod.main(["--writeback-dir", str(chosen)]) == 0
    assert (chosen / "runs").exists()
    assert not (ignored / "runs" / "run-").exists()


def test_explicit_paths_beat_the_working_area(fake_api, tmp_path, monkeypatch):
    import run as run_mod

    base = tmp_path / "area"
    src = _seed_inbox(base)
    elsewhere = tmp_path / "elsewhere"
    monkeypatch.setenv("PURVIEW_WRITEBACK_DIR", str(base))
    assert run_mod.main(["--workbook", str(src), "--out", str(elsewhere),
                         "--no-run-dir"]) == 0
    assert (elsewhere / "run_report.json").exists()
    assert not any((base / "runs").iterdir())


def test_a_tilde_path_is_expanded(monkeypatch):
    import run as run_mod

    monkeypatch.setenv("PURVIEW_WRITEBACK_DIR", "~/somewhere-under-home")
    resolved = run_mod.resolve_writeback_dir(None)
    assert resolved == Path.home() / "somewhere-under-home"
    assert "~" not in str(resolved)


def test_missing_folders_are_created(fake_api, tmp_path, monkeypatch, caplog):
    import run as run_mod

    base = tmp_path / "brand-new"
    monkeypatch.setenv("PURVIEW_WRITEBACK_DIR", str(base))
    run_mod.main([])            # no export present yet -- exits non-zero, by design
    assert (base / "inbox").is_dir()
    assert (base / "runs").is_dir()
    assert "ExportActions.xlsx" in caplog.text   # says exactly where to put it


def test_without_the_setting_nothing_changes(fake_api, tmp_path, monkeypatch):
    import run as run_mod

    monkeypatch.delenv("PURVIEW_WRITEBACK_DIR", raising=False)
    workbook = find_workbook()
    if workbook is None:
        pytest.skip("client export not present on this machine")
    assert run_mod.main(["--workbook", str(workbook), "--out", str(tmp_path),
                         "--no-run-dir"]) == 0
    assert (tmp_path / "run_report.json").exists()


# --- confirming an upload landed ------------------------------------------- #

def test_artifacts_can_be_looked_up_by_reference_id(cleanup_api):
    # Nobody has the UUID to hand; the reference id is what the summary prints.
    import upload as upload_mod

    evidence_id, artifacts = upload_mod.list_artifacts("PURVIEW-ACTION-UPDATE")
    assert evidence_id == "ev-known"
    assert len(artifacts) == 3


def test_a_uuid_is_used_directly_without_a_lookup(cleanup_api):
    import upload as upload_mod

    evidence_id, _ = upload_mod.list_artifacts("79edc659-f012-4c88-967c-be4c605c18f9")
    assert evidence_id == "79edc659-f012-4c88-967c-be4c605c18f9"


def test_an_unknown_reference_id_reports_nothing_uploaded(cleanup_api):
    import upload as upload_mod

    evidence_id, artifacts = upload_mod.list_artifacts("NEVER-UPLOADED")
    assert evidence_id is None and artifacts == []


def test_the_summary_always_says_whether_anything_was_uploaded(fake_api, tmp_path):
    import json as _json
    import report as report_mod

    assert _run(tmp_path) == 0
    report = _json.loads((tmp_path / "run_report.json").read_text())
    assert "NOT uploaded" in report_mod.summarize(report)

    assert _run(tmp_path, "--upload", "--min-match-rate", "0") == 0
    report = _json.loads((tmp_path / "run_report.json").read_text())
    text = report_mod.summarize(report)
    assert "UPLOADED" in text and "2 artifact(s)" in text


def test_brief_output_does_not_claim_nothing_was_written(fake_api, tmp_path):
    # Regression: --print-changes-brief hides Implementation Notes, and a run
    # that wrote only notes reported "(no cells written)" for 327 real writes.
    import report as report_mod

    report = {
        "totals": {"cells_written": 327},
        "rows": [{"action_name": "A", "updates": [
            {"column": "Implementation Notes", "old": None, "new": "x", "written": True}]}],
    }
    brief = report_mod.changes_table(report, skip_columns=("Implementation Notes",))
    assert "no cells written" not in brief
    assert "327 cell(s) written" in brief and "--print-changes" in brief
